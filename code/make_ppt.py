"""Builds presentation/Chess_Arena_Minimax_vs_AlphaBeta.pptx from results/results.json and figures/.
Run experiments.py first.      python make_ppt.py
"""
import json, os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.join(HERE, "..")
FIG = lambda n: os.path.join(ROOT, "figures", n)
R = json.load(open(os.path.join(ROOT, "results", "results.json")))
NAVY, BLUE, ORANGE, GREY, WHITE = (RGBColor(0x1b, 0x2a, 0x4a), RGBColor(0x2a, 0x6f, 0xdb),
                                   RGBColor(0xe0, 0x55, 0x2b), RGBColor(0x55, 0x55, 0x55), RGBColor(255, 255, 255))

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]


def slide(title, subtitle=None):
    s = prs.slides.add_slide(BLANK)
    bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(1.0))
    bar.fill.solid(); bar.fill.fore_color.rgb = NAVY; bar.line.fill.background()
    tf = bar.text_frame; tf.margin_left = Inches(0.5)
    p = tf.paragraphs[0]; p.text = title; p.font.size = Pt(28); p.font.bold = True; p.font.color.rgb = WHITE
    if subtitle:
        text(s, subtitle, 0.5, 1.1, 12.3, 0.5, 16, GREY, italic=True)
    return s


def text(s, t, x, y, w, h, size=18, color=None, bold=False, italic=False):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; p.text = t; p.font.size = Pt(size); p.font.bold = bold; p.font.italic = italic
    if color:
        p.font.color.rgb = color
    return tb


def bullets(s, items, x=0.6, y=1.4, w=12.1, h=5.6, size=20):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True
    for i, it in enumerate(items):
        sub = it.startswith("  ")
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = ("–  " if sub else "•  ") + it.strip()
        p.level = 1 if sub else 0
        p.font.size = Pt(size - 3 if sub else size)
        p.space_after = Pt(6)
    return tb


def table(s, rows, x, y, w, col_w=None, size=13):
    t = s.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w), Inches(0.4 * len(rows))).table
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            c = t.cell(i, j); c.text = str(v)
            p = c.text_frame.paragraphs[0]; p.font.size = Pt(size); p.font.bold = i == 0
            if i == 0:
                c.fill.solid(); c.fill.fore_color.rgb = BLUE; p.font.color.rgb = WHITE
    if col_w:
        for j, cw in enumerate(col_w):
            t.columns[j].width = Inches(cw)
    return t


def pic(s, name, x, y, w=None, h=None):
    kw = {}
    if w: kw["width"] = Inches(w)
    if h: kw["height"] = Inches(h)
    return s.shapes.add_picture(FIG(name), Inches(x), Inches(y), **kw)


nodes = R["nodes"]
dmax = max(r["depth"] for r in nodes)
at = lambda pos, d: next(r for r in nodes if r["position"] == pos and r["depth"] == d)
cases = R["cases"]
pick = lambda key, algo, d: next((r for r in cases[key]["by_depth"] if r["algo"] == algo and r["depth"] == d), None)


def arena_row(games, ab):
    w_ab = sum(1 for g in games if (g["result"] == "1-0" and g["white"] == ab) or (g["result"] == "0-1" and g["black"] == ab))
    w_mm = sum(1 for g in games if g["result"] in ("1-0", "0-1")) - w_ab
    ms = lambda algo: [m for g in games for m in g["moves"] if m["algo"] == algo]
    avg = lambda algo: sum(m["nodes"] for m in ms(algo)) / len(ms(algo))
    return w_ab, len(games) - w_ab - w_mm, w_mm, avg("minimax"), avg("alphabeta")


# 1 title
s = prs.slides.add_slide(BLANK)
bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
bg.fill.solid(); bg.fill.fore_color.rgb = NAVY; bg.line.fill.background()
text(s, "Multi-Agent Chess Arena", 0.8, 2.0, 11.5, 1.2, 48, WHITE, bold=True)
text(s, "A Minimax Agent vs. an Alpha-Beta Pruning Agent with Heuristic Evaluation", 0.8, 3.2, 11.5, 0.8, 24, WHITE)
text(s, "23CSE401 Fundamentals of Artificial Intelligence · Case Study", 0.8, 4.4, 11.5, 0.6, 18, RGBColor(0xbb, 0xcc, 0xee))
text(s, "Name: ____________________     Roll No.: ______________", 0.8, 5.6, 11.5, 0.6, 18, RGBColor(0xbb, 0xcc, 0xee))

# 2 problem
s = slide("Problem Statement")
bullets(s, ["Two AI agents play full chess games against each other in an arena",
            "  Agent 1: Minimax, which looks at every move sequence up to depth d",
            "  Agent 2: Alpha-Beta, the same search but it skips branches that cannot matter",
            "Both use the SAME heuristic evaluation (material + piece-square tables)",
            "Each agent's goal: maximise its own outcome (win > draw > loss); one agent's gain is the other's loss",
            "Interaction: strictly alternating moves on a shared board; every move changes the opponent's options",
            "Questions: Do both agents choose equally good moves? How much work does pruning save? "
            "What does that saving buy in actual games?"])

# 3 PEAS + environment
s = slide("PEAS and Environment Properties")
table(s, [["PEAS", "Chess arena"],
          ["Performance", "Win / draw / loss; material balance; nodes and time per move"],
          ["Environment", "8×8 board, 32 pieces, the opponent agent, rules of chess (python-chess)"],
          ["Actuators", "Choose one legal move per turn (incl. castling, en passant, promotion)"],
          ["Sensors", "Full board state (FEN), side to move, legal moves, check / mate / draw status"]],
      0.5, 1.4, 6.3, [1.6, 4.7], 13)
table(s, [["Property", "Value"],
          ["Observable", "Fully"],
          ["Deterministic", "Yes (no dice, no hidden cards)"],
          ["Episodic?", "Sequential"],
          ["Static?", "Static (turn-based)"],
          ["Discrete?", "Discrete"],
          ["Agents", "Multi-agent, competitive, zero-sum"],
          ["Known?", "Known rules"]],
      7.1, 1.4, 5.8, [2.0, 3.8], 13)

# 4 agent type
s = slide("Type of Agent")
bullets(s, ["Utility-based, model-based agent with adversarial look-ahead",
            "  Model: the rules of chess predict the result of every move (transition model)",
            "  Utility: the heuristic evaluation estimates how good a position is for White",
            "  Goal-based alone is not enough: checkmate is usually too far away to see, so positions must be scored",
            "Competitive, zero-sum: White maximises the evaluation and Black minimises it",
            "  Each agent assumes the opponent plays its best reply, which is exactly what minimax models",
            "Why not a reflex agent? Grabbing material greedily walks into traps (E1, E2)",
            "Rational within bounded depth: only as good as its search horizon"])

# 5 heuristic
s = slide("Heuristic Evaluation Function")
bullets(s, ["eval(s) = Σ White (piece value + square bonus) − Σ Black (piece value + square bonus)",
            "Piece values (centipawns): P 100, N 320, B 330, R 500, Q 900",
            "Piece-square tables reward central knights, advanced pawns, a castled king, …",
            "Terminal positions are scored exactly:",
            "  Checkmate = ±(100000 − ply), so a faster mate scores higher",
            "  Stalemate and insufficient material = 0",
            "Same function for both agents, so any difference comes from the SEARCH alone"], w=12.1)

# 6 algorithms
s = slide("Search Techniques")
text(s, "Minimax (depth-limited)", 0.6, 1.3, 6, 0.5, 22, BLUE, bold=True)
bullets(s, ["Depth-first over the game tree to depth d",
            "MAX (White) takes the max of its children and MIN (Black) the min",
            "Leaves are scored with the heuristic",
            "Visits every node: O(b^d) time, O(b·d) memory"], 0.6, 1.9, 6, 4, 17)
text(s, "Alpha-Beta pruning", 6.9, 1.3, 6, 0.5, 22, ORANGE, bold=True)
bullets(s, ["Carries α (best for MAX so far) and β (best for MIN so far)",
            "If α ≥ β, the opponent will never allow this line, so the remaining moves are skipped",
            "Returns EXACTLY the minimax value",
            "Move ordering (captures first, MVV-LVA) finds cut-offs early",
            "Best case O(b^(d/2)), worst case O(b^d)"], 6.9, 1.9, 6, 4, 17)

# 7 tree
s = slide("Search Tree: Where Pruning Happens", "King+pawn endgame, depth 2, same move order for both searches")
pic(s, "search_tree.png", 0.4, 1.7, w=12.5)
t = R["tree"]
text(s, f"Minimax visits {t['minimax_nodes']} positions; alpha-beta visits {t['alphabeta_nodes']}; both return value {t['value']}.",
     0.6, 6.7, 12, 0.5, 16, GREY)

# 8 nodes
s = slide("Nodes Searched vs. Depth", "Alpha-beta always returned the same value as minimax (checked by assert)")
pic(s, "nodes_vs_depth.png", 0.3, 1.6, w=8.2)
rows = [["Position", "Minimax", "Alpha-beta", "Saved"]]
for pos in dict.fromkeys(r["position"] for r in nodes):
    r = at(pos, dmax)
    rows.append([pos, f"{r['minimax']['nodes']:,}", f"{r['ab_mvv']['nodes']:,}",
                 f"{100 * (1 - r['ab_mvv']['nodes'] / r['minimax']['nodes']):.1f}%"])
table(s, rows, 8.6, 1.8, 4.5, [1.8, 1.0, 1.0, 0.7], 11)
text(s, f"Depth {dmax}", 8.6, 1.4, 4, 0.4, 14, GREY, bold=True)

# 9 working cases
s = slide("Working Cases")
pic(s, "case_W1.png", 0.4, 1.2, h=4.3)
pic(s, "case_W2.png", 6.9, 1.2, h=4.3)
w1a, w1m = pick("W1", "alphabeta", 3), pick("W1", "minimax", 3)
w2a, w2m = pick("W2", "alphabeta", 3), pick("W2", "minimax", 3)
text(s, f"W1 Back-rank mate: both agents play {w1a['move']} at every depth; the mate score is exact. "
        f"At depth 3: {w1m['nodes']:,} vs {w1a['nodes']:,} nodes.", 0.4, 5.6, 6.1, 1.5, 15)
text(s, f"W2 Knight fork: at depth 1-2 the fork is invisible (the rook capture is on ply 3). At depth 3 both play "
        f"{w2a['move']} (value {w2a['value']:+d}). Nodes: {w2m['nodes']:,} vs {w2a['nodes']:,}.", 6.9, 5.6, 6.1, 1.5, 15)

# 10 edge cases E1, E2
s = slide("Edge Cases: Horizon Effect and Stalemate Trap")
pic(s, "case_E1.png", 0.4, 1.2, h=4.3)
pic(s, "case_E2.png", 6.9, 1.2, h=4.3)
e1s = [r for r in cases["E1"]["by_depth"] if r["algo"] == "alphabeta"]
text(s, "E1 (position from an agent game): at depth 1-2 Black grabs a pawn with " + e1s[0]["move"] +
        ", but Bf4 then pins the queen to the king and wins it two plies later. At depth ≥ 3 the agent sees this and plays "
        + e1s[-1]["move"] + ".", 0.4, 5.6, 6.2, 1.6, 14)
text(s, f"E2: Qxe6 wins a knight but is STALEMATE (value 0). A greedy material-counter plays "
        f"{cases['E2']['greedy_static']}; both search agents play {pick('E2', 'alphabeta', 3)['move']} and find a forced mate.",
     6.9, 5.6, 6.1, 1.6, 14)

# 11 move ordering edge case
s = slide("Edge Case: Bad Move Ordering", "Alpha-beta's saving depends on examining good moves first")
r = at("Italian opening", dmax)
table(s, [["Italian opening, depth " + str(dmax), "Nodes", "vs minimax"],
          ["Minimax", f"{r['minimax']['nodes']:,}", "100%"],
          ["Alpha-beta, reversed ordering (bad moves first)", f"{r['ab_worst']['nodes']:,}", f"{100 * r['ab_worst']['nodes'] / r['minimax']['nodes']:.1f}%"],
          ["Alpha-beta, generator order", f"{r['ab_none']['nodes']:,}", f"{100 * r['ab_none']['nodes'] / r['minimax']['nodes']:.1f}%"],
          ["Alpha-beta, MVV-LVA ordering", f"{r['ab_mvv']['nodes']:,}", f"{100 * r['ab_mvv']['nodes'] / r['minimax']['nodes']:.1f}%"]],
      0.8, 1.9, 11.5, [6.5, 2.5, 2.5], 16)
bullets(s, ["Same value in every run: ordering changes only the work, never the answer",
            "With good moves first, one refutation is enough to prune the rest",
            "With bad moves first, cut-offs come late and alpha-beta approaches minimax, towards O(b^d)"], 0.8, 4.7, 11.5, 2.5, 18)

# 12 arena
s = slide("Arena Results", "6 openings × both colours, 100-ply limit")
pic(s, "arena.png", 0.3, 1.6, w=7.6)
ra, rb = arena_row(R["arena_equal_depth"], "alphabeta:2"), arena_row(R["arena_deeper_ab"], "alphabeta:3")
table(s, [["Match", "AB W-D-L", "MM nodes/move", "AB nodes/move"],
          ["A: MM d2 vs AB d2", f"{ra[0]}-{ra[1]}-{ra[2]}", f"{ra[3]:,.0f}", f"{ra[4]:,.0f}"],
          ["B: MM d2 vs AB d3", f"{rb[0]}-{rb[1]}-{rb[2]}", f"{rb[3]:,.0f}", f"{rb[4]:,.0f}"]],
      8.1, 1.8, 5.0, [1.9, 1.0, 1.05, 1.05], 12)
ag = R["agreement"]
bullets(s, [f"Equal depth: in {ag['positions']} replayed positions both agents' moves had the SAME value "
            f"({100 * ag['same_value'] / ag['positions']:.0f}%) and were identical in {100 * ag['same_move'] / ag['positions']:.1f}%; "
            "the score comes only from tie-breaks",
            "Alpha-beta at depth 3 costs less than 2× minimax's nodes at depth 2, and the extra ply wins games"],
        8.1, 3.4, 5.0, 3.5, 15)

# 13 complexity
s = slide("Complexity")
st = at("Start position", dmax)
table(s, [["", "Time", "Space", "Chess (b ≈ 35)"],
          ["Minimax", "O(b^d)", "O(b·d)", "d=4: ≈ 1.5 million leaves"],
          ["Alpha-beta, best case", "O(b^(d/2))", "O(b·d)", "d=4: ≈ 2,450 leaves; twice the depth for the same work"],
          ["Alpha-beta, worst case", "O(b^d)", "O(b·d)", "no better than minimax"]],
      0.6, 1.4, 12.1, [2.8, 1.6, 1.3, 6.4], 15)
bullets(s, [f"Measured, start position (b = 20), depth {dmax}: minimax {st['minimax']['nodes']:,} nodes, "
            f"alpha-beta {st['ab_mvv']['nodes']:,}",
            "Each extra ply multiplies minimax's work by about b (20-40×), and alpha-beta's by about √b",
            "Full chess tree ≈ 10^120 nodes (Shannon), so any agent must cut off at depth d and use a heuristic"],
        0.6, 3.6, 12.1, 3.5, 18)

# 14 comparison
s = slide("Comparison and Key Observations")
table(s, [["", "Minimax", "Alpha-Beta"],
          ["Decision quality at equal depth", "Optimal w.r.t. heuristic", "Identical value"],
          ["Nodes at depth " + str(dmax) + " (avg of 5 positions)",
           f"{sum(at(p, dmax)['minimax']['nodes'] for p in dict.fromkeys(r['position'] for r in nodes)) / 5:,.0f}",
           f"{sum(at(p, dmax)['ab_mvv']['nodes'] for p in dict.fromkeys(r['position'] for r in nodes)) / 5:,.0f}"],
          ["Sensitive to move ordering?", "No", "Yes (E3)"],
          ["Horizon effect?", "Yes", "Yes, but it can search deeper and so push the horizon back"],
          ["Practical depth in Python (~1 s/move)", "≈ 3", "≈ 4-5"]],
      0.6, 1.4, 12.1, [4.3, 3.6, 4.2], 15)
bullets(s, ["Pruning is free accuracy-wise: it removes work, not information",
            "Depth matters more than anything else: the deeper agent avoids traps (E1) and finds tactics (W2)",
            "Limitations: horizon effect, a heuristic blind to king safety and mobility, exponential growth"],
        0.6, 4.6, 12.1, 2.5, 18)

# 15 conclusion
s = slide("Conclusion")
bullets(s, ["Chess is a fully observable, deterministic, sequential, competitive two-agent environment, which is ideal for adversarial search",
            "Minimax and alpha-beta agents with the same evaluation choose moves of the same value; alpha-beta does it with "
            "1-2 orders of magnitude less work",
            "Good move ordering is what makes the pruning effective",
            "In the arena, the saved effort becomes extra depth, and extra depth becomes wins",
            "Next steps: iterative deepening, quiescence search (fixes the horizon effect), transposition tables"])

out = os.path.join(ROOT, "presentation", "Chess_Arena_Minimax_vs_AlphaBeta.pptx")
os.makedirs(os.path.dirname(out), exist_ok=True)
prs.save(out)
print("saved", out)
